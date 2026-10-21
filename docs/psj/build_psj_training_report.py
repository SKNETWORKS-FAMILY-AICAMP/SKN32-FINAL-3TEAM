from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"C:\SKN32-FINAL-3TEAM\docs\psj\머신러닝_딥러닝_학습결과서_2차_박수진_제출본.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.top_margin = sec.bottom_margin = Cm(2.05)
sec.left_margin = sec.right_margin = Cm(2.25)

styles = doc.styles
normal = styles['Normal']
normal.font.name = '맑은 고딕'
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor(25, 33, 43)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.16
for sty, size, before, after in [('Title', 19, 0, 8), ('Heading 1', 12, 15, 7), ('Heading 2', 10.5, 10, 5)]:
    s = styles[sty]
    s.font.name = '맑은 고딕'
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    if sty == 'Title':
        pPr = s.element.get_or_add_pPr()
        for child in list(pPr):
            if child.tag == qn('w:pBdr'):
                pPr.remove(child)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        x=OxmlElement('w:'+edge)
        x.set(qn('w:val'),'single')
        x.set(qn('w:sz'),'4')
        x.set(qn('w:color'),'D9D9D9')
        b.append(x)
    tcPr.append(b)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for i, x in enumerate(headers): t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            border(c)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: c.width=Cm(widths[ci])
            if ri==0: shade(c,'DDE8F3')
            elif ri%2==0: shade(c,'F7F9FB')
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.space_after=Pt(3)
                if ci>0 and len(headers)>2: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name='맑은 고딕'
                    run.font.size=Pt(8.2)
                    if ri==0: run.font.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def p(text='', boldlead=None):
    para=doc.add_paragraph()
    if boldlead:
        r=para.add_run(boldlead)
        r.bold=True
    para.add_run(text)
    return para

doc.add_paragraph('판정 인코더 머신러닝 딥러닝 학습결과서 2차', 'Title')
p('CopyLane  |  박수진  |  제출일 2026년 9월 28일  |  실험 기준일 2026년 9월 22일')

doc.add_heading('1 결과 요약',1)
p('판정 인코더의 8개 위법 유형 다중 라벨 분류 실험에서 beomi/kcbert-base와 Focal Loss γ=1.5, 양성 가중치 상한 5를 최종 설정으로 채택했다. 9월 22일 평가 문구 352건에서 기본 임계값 0.5 기준 micro F1은 0.444였다. γ=1.5·상한 1의 0.409, γ=1.5·상한 10의 0.430보다 높았다. 무해 문구 60건의 오판정은 이 실행에서 0건이었다.')
p('이 수치는 9월 22일 데이터와 유형 단위 평가 결과다. 9월 25일 골든셋은 8,432행으로 개정되고 정답 기준이 조문 인용·호 단위로 바뀌었다. 개정 데이터에서 이 모델의 성능은 아직 측정되지 않았으므로, 본 결과를 최신 평가셋의 성능이나 전체 배포 통과로 해석하지 않는다.')

doc.add_heading('2 학습 대상과 실행 조건',1)
p('입력 광고 문구에 대해 8개 위법 유형을 각각 예측하는 다중 라벨 분류기다. 문서 전체의 위험도 조합과 조문 인용 생성은 이번 실험 범위에 포함하지 않았다. 학습·평가 분할은 당시 입력 파일의 split 필드를 그대로 사용했다.')
table(['항목','9월 22일 실험 설정'],[
('학습 데이터','golden.jsonl 6,741행 + injected_golden.jsonl 902행, 합계 7,643행'),
('분할','train 7,291행 / 평가 352행(위반 292, 무해 60)'),
('백본과 출력','beomi/kcbert-base / 위법 유형 8개 다중 라벨'),
('손실과 채택값','Focal Loss, γ=1.5, pos_weight 상한 5'),
('학습 조건','3 epoch, 학습률 2×10⁻⁵, 최대 길이 128, 배치 16'),
('평가 기준','기본 임계값 0.5; micro F1 및 유형별 Precision·Recall·F1')],[3.1,13.7])
p('입력 파일의 해시와 각 실행의 체크포인트 식별자는 당시 결과 문서에 남아 있지 않다. 따라서 본 보고서는 제공된 실험 문서와 노트북의 기록을 재구성한 결과이며, 동일 데이터의 독립 재실행 증빙은 아니다.')

h3=doc.add_heading('3 설정 비교와 채택 근거',1)
h3.paragraph_format.page_break_before=True
p('먼저 pos_weight 상한 1에서 γ를 비교하고, γ=1.5를 고정한 다음 상한 5와 10을 비교했다. 상한 5에서 최고 micro F1을 얻었고, 상한 10은 성능과 무해 문구 오판정률 모두 악화됐다.')
table(['γ','상한','micro F1','무해 오판정','F1 0 유형'],[
('1.0','1','0.394','1.7%','4'),('1.5','1','0.409','0.0%','4'),
('2.0','1','0.406','1.7%','4'),('3.0','1','0.387','0.0%','4'),
('1.5','5','0.444','0.0%','3'),('1.5','10','0.430','1.7%','3')],[2.3,2.1,3.5,4.6,4.3])

doc.add_heading('4 채택 설정의 유형별 결과',1)
p('다음 표는 동일 평가셋에서 임계값 0.5를 적용한 결과다. 양성 표본이 30건 미만인 유형은 변동성이 커 참고값으로만 본다.')
table(['위법 유형','평가 양성','F1','해석'],[
('거짓 과장','108','0.548','상대 Recall 기준 통과'),('소비자 기만','70','0.502','상대 Recall 기준 통과'),
('건강기능식품 오인','53','0.000','미달'),('부당 비교광고','35','0.103','미달'),
('비방광고','12','0.275','표본 부족'),('의약품 오인','8','0.000','표본 부족'),
('질병 예방치료 표방','6','0.250','표본 부족'),('후기 체험기 기만','1','0.000','표본 부족')],[5.0,2.5,2.6,6.7])
p('기준선 B보다 Recall을 0.15 이상 높인다는 당시 상대 기준에서는 거짓 과장과 소비자 기만만 통과했다. 건강기능식품 오인과 부당 비교광고는 미달했다. 따라서 8개 유형 전체가 배포 기준을 충족했다고 볼 수 없다.')

h5=doc.add_heading('5 임계값 탐색과 해석상의 제한',1)
h5.paragraph_format.page_break_before=True
p('평가셋에서 라벨별 임계값을 탐색했을 때 micro F1은 0.447, Precision은 0.320, Recall은 0.741이었다. 기본 임계값의 micro F1 0.444와 비교하면 재현율은 높아졌지만 정밀도는 낮아졌다. 같은 평가셋에서 임계값을 선택하고 성능을 재면 낙관 편향이 생기므로 이 수치는 탐색 결과로만 기록한다. 별도 검증셋에서 임계값을 결정한 뒤 봉인 평가셋에서 다시 측정해야 한다.')
p('무해 오판정 0건 역시 60건의 표본에서 관찰된 값이다. 의약품 오인은 임계값을 조정해도 F1 0이었으며, 건강기능식품 오인과 후기 체험기 기만은 기본 임계값에서 F1 0이었다. 데이터 보강 없이 손실함수 조정만으로 해결됐다고 보기 어렵다.')

doc.add_heading('6 최신 데이터와 재검증 계획',1)
p('9월 25일판 골든셋은 총 8,432행(train 6,425, test_sentence 2,007)이다. 최신 정답은 조문 인용이며 유형명은 파생값이다. 채점은 호 단위로 하고 조건 D·M을 제외한 1,555행을 대상으로 한다. 이는 9월 22일의 유형 단위 352행 평가와 다른 과제이므로 두 수치를 직접 비교할 수 없다.')
p('최신 golden.jsonl을 공유 저장소에서 동기화하고 split_manifest.json과 doc_id 단위 분할을 확인한다. 이후 채택한 γ=1.5·상한 5를 출발 설정으로 재학습한다. train 내부 검증에서 임계값을 정한 뒤 봉인 평가셋을 한 번만 채점하고, 호별 Precision·Recall·F1, 30건 미만 호의 측정 불가, 적법 60행 오탐을 보고한다. 해설서는 전량 평가에만 두고 주입 행을 다시 더하지 않는다.')

doc.add_heading('7 산출물과 근거 자료',1)
p('제공된 모델 파일: copylane-encoder-kcbert-final-cap5-gamma1.5.zip. 가중치, tokenizer, config, label_scheme.json이 포함돼 있다. 다만 ZIP 내부 메타데이터만으로 γ와 가중치 상한을 독립 확인할 수 없으므로 최종 설정의 근거는 실험 문서와 실행 기록이다. 9월 22일 모델과 9월 25일 개정 데이터의 모델은 별개로 관리해야 한다.')
p('근거: 판정_인코더_파인튜닝_결과_0922최종_박수진.docx; copylane_encoder_finetune_v3_gamma_sweap.ipynb; copylane_encoder_finetune_v4_0_pos_weight_cap_sweap (1).ipynb; docs/03_데이터/전처리_결과서_2차.md; docs/03_데이터/학습결과서_2차.md; 설계결정기록 D-70·D-174·D-231·D-282·D-286.')

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('CopyLane  ·  박수진  ·  2026.09.28')
r.font.name='맑은 고딕'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,110,120)

doc.save(OUT)
print(OUT)
